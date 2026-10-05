from docx import Document

from core.rag.parser import extract_text_from_path


def test_docx_specification_tables_remain_between_their_context_paragraphs(tmp_path):
    doc = Document()
    doc.add_paragraph('虚构产品 Acceptance_widget')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'MOQ'
    table.cell(0, 1).text = '73 units'
    table.cell(1, 0).text = '交期'
    table.cell(1, 1).text = '19 days'
    doc.add_paragraph('仅供验收')
    path = tmp_path / 'specs.docx'
    doc.save(path)

    assert extract_text_from_path(str(path)).splitlines() == [
        '虚构产品 Acceptance_widget', 'MOQ\t73 units', '交期\t19 days', '仅供验收',
    ]
