"""Real DOCX/PDF parsing, indexing and tenant boundaries in the isolated fixture."""
import io
import json
from pathlib import Path


def document_fixtures():
    from docx import Document
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

    word = Document()
    word.add_paragraph('Fictional Acceptance_widget specifications')
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'MOQ'
    table.cell(0, 1).text = '73 units'
    table.cell(1, 0).text = 'Lead time'
    table.cell(1, 1).text = '19 days'
    word.add_paragraph('End of fictional specifications.')
    word_bytes = io.BytesIO()
    word.save(word_bytes)
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                             NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 14 Tf 40 780 Td (Fictional Acceptance_widget MOQ 73 units; lead time 19 days.) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    pdf_bytes = io.BytesIO()
    writer.write(pdf_bytes)
    return [('acceptance-table.docx', word_bytes.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
            ('acceptance-specs.pdf', pdf_bytes.getvalue(), 'application/pdf')]


def verify_documents(client, checked, factory, args, kb_id):
    from database.shared_models import KnowledgeDoc, KnowledgeChunk
    from core.rag.retriever import KnowledgeRetriever

    rows = []
    for name, data, mime in document_fixtures():
        if args.evidence_dir:
            args.evidence_dir.mkdir(parents=True, exist_ok=True)
            (args.evidence_dir / name).write_bytes(data)
        uploaded = checked(client.post(f'/api/v1/kb/bases/{kb_id}/docs', files={'file': (name, data, mime)}))
        with factory() as db:
            doc = db.get(KnowledgeDoc, uploaded['id'])
            chunks = db.query(KnowledgeChunk).filter_by(doc_id=doc.id).all()
            assert doc.status == ('embedded' if args.external_services else 'indexed'), f'{name}: indexing failed'
            text = '\n'.join(chunk.chunk_text for chunk in chunks)
            assert '73' in text and '19' in text, f'{name}: document facts lost during parsing'
            count = len(chunks)
            assert count > 0
            if args.external_services:
                assert all(chunk.embedding is not None and len(chunk.embedding) == 1024 for chunk in chunks)
        checked(client.post(f"/api/v1/kb/docs/{uploaded['id']}/reindex"))
        with factory() as db:
            assert db.query(KnowledgeChunk).filter_by(doc_id=uploaded['id']).count() == count, 'Reindex duplicated chunks'
            assert db.get(KnowledgeDoc, uploaded['id']).status == ('embedded' if args.external_services else 'indexed')
        rows.append({'filename': name, 'id': uploaded['id'], 'chunks': count, 'reindex_without_duplicates': True})

    with factory() as db:
        if args.external_services:
            retriever = KnowledgeRetriever(db)
            query = '最小采购量和生产周期分别是多少？'
            assert retriever._lexical_search([kb_id], query, 5) == [], 'Semantic proof must not rely on lexical matching'
            hits = retriever._vector_search([kb_id], query, 5)
            names = {hit['doc_name'] for hit in hits if '73' in hit['content'] and '19' in hit['content']}
            assert {row['filename'] for row in rows}.issubset(names), 'Cross-language vector retrieval missed a format'

    bad = checked(client.post(f'/api/v1/kb/bases/{kb_id}/docs', files={'file': ('invalid.pdf', b'not a PDF', 'application/pdf')}))
    failed = next(d for d in checked(client.get(f'/api/v1/kb/bases/{kb_id}/docs'))['items'] if d['id'] == bad['id'])
    assert failed['status'] == 'failed' and failed['error_msg'] and failed['chunk_count'] == 0
    owner_auth = client.headers['Authorization']
    outsider = checked(client.post('/auth/register', json={'email': 'other-docs@example.invalid', 'password': 'Local-Acceptance-only!2026'}), 201)
    client.headers['Authorization'] = 'Bearer ' + outsider['access_token']
    other_org = checked(client.post('/auth/organization/create', json={'name': 'Other document tenant'}))
    client.headers['Authorization'] = 'Bearer ' + other_org['access_token']
    try:
        assert client.get(f'/api/v1/kb/bases/{kb_id}/docs').status_code == 403
        assert client.post('/api/v1/kb/search', json={'query': 'MOQ', 'kb_ids': [kb_id]}).status_code == 403
        assert client.post(f"/api/v1/kb/docs/{rows[0]['id']}/reindex").status_code == 403
        assert client.delete(f"/api/v1/kb/docs/{rows[0]['id']}").status_code == 403
    finally:
        client.headers['Authorization'] = owner_auth
    with factory() as db:
        stored = Path(db.get(KnowledgeDoc, bad['id']).file_path)
    checked(client.delete(f"/api/v1/kb/docs/{bad['id']}"))
    assert not stored.exists()
    receipt = {'documents': 'passed', 'formats': rows, 'word_table_facts_preserved': True,
               'semantic_cross_language_verified': args.external_services, 'invalid_pdf_reports_failure': True,
               'cross_tenant_read_search_reindex_delete_denied': True, 'deleted_upload_removed': True, 'mock_used': False}
    if args.evidence_dir:
        (args.evidence_dir / 'knowledge-documents-receipt.json').write_text(json.dumps(receipt, indent=2), encoding='utf-8')
    print(json.dumps(receipt), flush=True)
